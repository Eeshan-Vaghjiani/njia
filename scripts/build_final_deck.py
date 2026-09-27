"""Local-only Njia deck builder. No API calls, credentials, or remote assets.

Everything:  .venv/bin/python scripts/build_final_deck.py --all [--out DIR]
Build only:  .venv/bin/python scripts/build_final_deck.py [--out DIR]
PDF + PNGs:  .venv/bin/python scripts/build_final_deck.py --render-only [--out DIR]
Validate:    .venv/bin/python scripts/build_final_deck.py --validate [--out DIR]
Writes only inside --out (default presentation/). PPTX and HTML share the same design objects.
Every tested or observed figure comes from EVIDENCE below or the bundled dataset.
"""
from pathlib import Path
import argparse
import collections
import html
import json
import math
import re

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation"
STEM = "Njia-Dhruzzz"
DEMO_URL = "https://github.com/Eeshan-Vaghjiani/njia/releases/download/demo-v1/njia-demo-90s.webm"

# Single source for every tested/observed figure shown in the deck, notes and README.
# The first block carries over the verified values from the published deck. The
# new-feature block stays None until the integrator has verified a value; None never
# renders a number, only a neutral description of the feature.
EVIDENCE = {
    "backend_tests": 73,
    "backend_count_source": "Main implementation owner confirmation",
    "acceptance_assertions": 48,
    "recording_assertions": 9,
    "provider": "Groq",
    "model": "openai/gpt-oss-20b",
    "observed_seconds": {"desktop": 2.259, "mobile": 1.318, "recording_with_pdf": 2.260},
    "javascript_errors": 0,
    "sample_size": 391,
    "coverage": 39.6,
    "video_seconds": 90,
    "source": "Verified values carried over from the published deck (coach acceptance + demo recording runs).",
    "factual_limit": "Observed unsupported real-time sales monitoring in a rewrite; human factual review required.",
    # New features. Fill only with verified values, e.g.
    # "live_jobs_example": {"count": 111, "role": "Data Analyst", "country": "Kenya", "date": "27 Sep 2026"},
    # "web_questions_example": {"count": 5, "role": "Data Analyst", "country": "Kenya"},
    # "new_feature_tests": 24,
    "live_jobs_example": None,
    "web_questions_example": None,
    "new_feature_tests": None,
    "demo_predates_new_features": True,
}
NEW_FEATURE_KEYS = ["live_jobs_example", "web_questions_example", "new_feature_tests"]
TAGLINE = "Ask. Match. Practise. Your next move."
SANS = "'Segoe UI','Noto Sans','DejaVu Sans',Arial,sans-serif"
SERIF = "Georgia,'Noto Serif','DejaVu Serif',serif"
NAVY = "14283F"
IVORY = "F7F4EC"
ORANGE = "F27B43"
INK = "1B2D40"
MUTED = "586777"
PALE = "D4DEE6"
LINE = "DDDCD5"
WHITE = "FFFFFF"
LIGHT = "FFF0E5"
GREEN = "246855"
SLIDES = []


def public_total():
    return EVIDENCE["acceptance_assertions"] + EVIDENCE["recording_assertions"]


def seconds(key):
    return f'{EVIDENCE["observed_seconds"][key]:.3f}s'


def new_feature_parts():
    jobs, questions, tests = (EVIDENCE[k] for k in NEW_FEATURE_KEYS)
    return [
        (f'{jobs["count"]:,} remote {jobs["role"]} jobs open to {jobs["country"]} ({jobs["date"]})' if jobs
         else "Remote jobs open to Kenya or worldwide"),
        (f'{questions["count"]} source-linked interview questions' if questions
         else "only source-verified web questions shown"),
        (f"{tests} new-feature tests passed" if tests else "labelled fallback on provider failure"),
    ]


def box(s, x, y, w, h, fill, radius=0, stroke=None):
    s["items"].append(dict(kind="box", x=x, y=y, w=w, h=h,
                           fill=fill, radius=radius, stroke=stroke))


def text(s, value, x, y, w, h, size=27, color=None, bold=False,
         serif=False, align="left", link=None):
    s["items"].append(dict(kind="text", text=value, x=x, y=y, w=w, h=h,
                           size=size, color=color or s["fg"], bold=bold,
                           serif=serif, align=align, link=link))


def slide(kicker, title, subtitle=None, dark=False, notes=""):
    n = len(SLIDES) + 1
    s = dict(items=[], dark=dark, bg=NAVY if dark else IVORY,
             fg=IVORY if dark else INK, notes=notes, title=title)
    SLIDES.append(s)
    box(s, 72, 54, 38, 5, ORANGE)
    text(s, kicker.upper(), 129, 40, 1320, 31, 19, ORANGE if dark else MUTED, True)
    text(s, title, 72, 109, 1456, 150, 65, bold=True, serif=True)
    if subtitle:
        text(s, subtitle, 75, 275, 1440, 70, 26, PALE if dark else MUTED)
    box(s, 72, 838, 1456, 1, "385065" if dark else LINE)
    text(s, "NJIA  /  DHRUZZZ  /  KENYA · ONLINE", 74, 853, 1100, 25, 16,
         PALE if dark else MUTED)
    text(s, f"{n:02d} / 08", 1400, 851, 128, 28, 18,
         PALE if dark else MUTED, align="right")
    return s


def label(s, value, x, y, w=600, color=ORANGE):
    text(s, value.upper(), x, y, w, 32, 19, color, True)


def build_content():
    e = EVIDENCE
    bt, acc, rec, total = e["backend_tests"], e["acceptance_assertions"], e["recording_assertions"], public_total()
    jobs = e["live_jobs_example"]
    jobs_note = (f' In a live check on {jobs["date"]}, {jobs["count"]:,} remote {jobs["role"]} postings were open to'
                 f' {jobs["country"]} or worldwide.' if jobs else "")
    rows = [json.loads(line) for line in
            (ROOT / "data/africa_jobs_subset.jsonl").read_text(encoding="utf-8").splitlines() if line]
    ke = [r for r in rows if r["job_country"] == "Kenya"]
    da = [r for r in ke if r["job_title_short"] == "Data Analyst"]
    countries = len({r["job_country"] for r in rows})
    assert (len(rows), countries, len(ke), len(da)) == (18371, 10, 1326, 391)
    assert len(da) == e["sample_size"]
    counts = collections.Counter(skill for r in da for skill in set(r["skills"]))
    skills = ["sql", "python", "r", "excel", "spss"]
    demand = sorted([(skill, round(counts[skill] / len(da) * 100, 1)) for skill in skills], key=lambda item: -item[1])
    (OUT / "production-evidence.json").write_text(json.dumps({
        "backend_tests_passed": bt, "backend_count_source": e["backend_count_source"],
        "public_acceptance_assertions": acc, "recording_assertions": rec, "public_total": total,
        "provider": e["provider"], "model": e["model"], "observed_seconds": e["observed_seconds"],
        "javascript_errors": e["javascript_errors"], "sample_size": e["sample_size"], "coverage": e["coverage"],
        "video_seconds": e["video_seconds"], "source": e["source"], "factual_limit": e["factual_limit"],
        **{k: e[k] for k in NEW_FEATURE_KEYS},
        "unverified_new_feature_fields": [k for k in NEW_FEATURE_KEYS if e[k] is None],
        "demo_predates_new_features": e["demo_predates_new_features"],
    }, indent=2), encoding="utf-8")

    s = slide("01  /  The Kenyan learner", "A CV tells your past.\nNjia helps choose your next step.", dark=True,
              notes="Njia means path in Swahili. We are Dhruzzz, a Kenya ONLINE team. Wanjiru is an illustrative persona, not an interviewed user. Njia now asks before it advises: after the CV is provided, it asks a few follow-up questions, then links strengths to CV evidence and the learner's own answers, sets three priorities from historical 2023 Kenya role data, and suggests fact-preserving CV rewrites and a seven-day plan. It also answers two new questions: which jobs can I apply for now, and am I ready for the interview? Skill-match percentages measure skill overlap and are not a hiring probability. Evidence and limits are on slide 6.")
    text(s, "A career coach that asks before it advises,\nthen matches you to jobs you can apply for.", 76, 295, 940, 105, 34, PALE)
    label(s, "THE LEARNER'S THREE QUESTIONS", 76, 461)
    for y, n, heading, sub in [
        (518, "01", "What can I already show?", "Njia asks follow-ups, then links strengths to CV evidence."),
        (612, "02", "Which jobs can I apply for now?", "Live postings with a skill-match %, not hiring odds."),
        (706, "03", "Am I ready for the interview?", "Real reported questions, typed practice and AI feedback.")]:
        text(s, n, 76, y, 55, 45, 27, ORANGE, True)
        text(s, heading, 148, y, 790, 42, 29, IVORY, True)
        text(s, sub, 148, y + 42, 860, 39, 23, PALE)
    box(s, 1050, 300, 478, 481, "203A52", 18)
    text(s, "W", 1092, 328, 130, 120, 90, ORANGE, serif=True)
    label(s, "ILLUSTRATIVE PERSONA", 1095, 465, 380, PALE)
    text(s, "Wanjiru\nNairobi · aspiring data analyst", 1095, 511, 388, 112, 31, IVORY, True)
    text(s, "Has project experience.\nUnsure which roles fit\nand how to prepare.", 1095, 649, 389, 107, 26, PALE)

    s = slide("02  /  The product story", "Ask first. Then match and practise.",
              "Follow-up questions, an evidence-linked brief, live job matches and interview practice.",
              notes="The product story is Ask, Match, Practise. Ask: after the CV is provided, the AI asks three to five follow-up questions: skill questions with the options Used it at work, Used it in a project or course, Still learning it, or Not yet, plus a free-text question about an outcome. Answers reshape the suggested skills and the brief. The before/after example is a hand-authored illustration, not a captured model response. The rewrite uses only the CV line and the learner's own answer; anything drawn from an answer is labelled Uses your answer — verify. No business impact, percentage or credential is invented. Match: live postings the learner can apply for, each with a skill-match percentage, matched and missing skills, and a link to apply on the source. Practise: interview questions candidates report on the web, each with its source link, then typed practice with AI feedback. The seven-day checklist can also be sent to WhatsApp.")
    for x, w, fill in [(74, 660, WHITE), (774, 754, NAVY)]:
        box(s, x, 375, w, 242, fill, 14)
    label(s, "BEFORE / CV LINE + YOUR ANSWER", 104, 400, 580, MUTED)
    text(s, '“Cleaned sales data in Excel\nand made weekly charts.”', 104, 445, 580, 84, 35, INK, serif=True)
    text(s, "Your answer: “Managers used them in\nMonday stock reviews.”", 104, 541, 580, 56, 22, MUTED)
    label(s, "AFTER / USES YOUR ANSWER — VERIFY", 808, 400, 680)
    text(s, '“Prepared weekly Excel sales charts\nfrom cleaned data for managers’\nMonday stock reviews.”', 808, 445, 690, 125, 35, IVORY, serif=True)
    text(s, "Same facts plus the learner's own answer. Nothing invented.", 808, 577, 690, 26, 19, PALE)
    for x, heading, detail in [
        (74, "01 / Ask", "Skill + outcome follow-ups\nreshape skills and the brief."),
        (571, "02 / Match", "Live jobs you can apply for,\nwith matched + missing skills."),
        (1068, "03 / Practise", "Sourced questions + AI feedback;\nsend the 7-day plan to WhatsApp.")]:
        box(s, x, 653, 455, 4, ORANGE)
        text(s, heading, x, 680, 455, 43, 28, INK, True)
        text(s, detail, x, 732, 460, 72, 24, MUTED)
    text(s, "Illustrative, hand-authored example. Rewrites are checked for numbers/tools and still need human factual review.", 76, 336, 1440, 28, 18, MUTED)

    s = slide("03  /  Local data, visible denominators", "Better priorities start with local context.",
              "Historical skill mentions inform the advice; the language model does not invent demand figures.",
              notes="The chart is recalculated directly from data/africa_jobs_subset.jsonl. The full subset has 18,371 postings across ten countries, including 1,326 Kenyan postings across roles. The denominator for this chart is 391 Kenya Data Analyst postings. A posting may mention several skills, so the bars are not parts of a total. Source: Hugging Face lukebarousse/data_jobs, Apache-2.0, 2023. This is a historical job-posting sample, not live hiring demand or a representative census of Kenya's labor market. It sets the three priorities in the brief. Live openings are a separate feature, shown on slide 5, and never change these historical figures.")
    for x, big, caption in [(76, "18,371", "2023 postings in the subset"),
                             (573, "10", "countries represented"),
                             (1070, "1,326", "Kenya postings · all roles")]:
        text(s, big, x, 364, 450, 82, 68, INK, True, True)
        text(s, caption, x, 453, 455, 35, 23, MUTED)
    label(s, f"KENYA / DATA ANALYST / n = {len(da)}", 76, 530, 850, INK)
    for i, (skill, pct) in enumerate(demand):
        y = 580 + i * 40
        text(s, {"sql": "SQL", "python": "Python", "r": "R", "excel": "Excel", "spss": "SPSS"}[skill],
             76, y - 2, 132, 35, 24, INK, True)
        box(s, 217, y + 1, 575, 23, "E7E6DF", 3)
        box(s, 217, y + 1, 575 * pct / 50, 23, ORANGE if skill == "sql" else NAVY, 3)
        text(s, f"{pct:.1f}%", 810, y - 4, 110, 36, 24, INK, True)
    box(s, 985, 527, 543, 242, LIGHT, 12)
    text(s, "A direction, not a prediction.", 1017, 555, 480, 84, 37, INK, True, True)
    text(s, "Share of sampled posts naming each skill.\nOne post can mention several skills.\n2023 sample; not current vacancies.", 1017, 653, 490, 93, 23, MUTED)
    text(s, "Source: Hugging Face · lukebarousse/data_jobs · Apache-2.0 · 2023", 76, 795, 1400, 26, 17, MUTED,
         link="https://huggingface.co/datasets/lukebarousse/data_jobs")

    s = slide("04  /  AI architecture & relevance", "Ground the facts. Personalize the next step.",
              "AI asks and coaches; code computes market figures and keeps only source-verified web results.", dark=True,
              notes=f"The brief uses Groq-hosted {e['model']}; public production desktop, mobile and recording runs of the earlier release confirmed this model. CV text goes to the AI provider only after opt-in and basic, best-effort PII redaction, which is not full anonymization. There is no application database persistence of the CV; this is not a provider-retention claim. New today: the AI asks three to five follow-up questions after the CV, and the answers reshape suggested skills and the brief. Remote jobs come from the Himalayas public jobs API, filtered to postings whose location restrictions include the user's country or worldwide. Local postings and reported interview questions come from Groq's built-in browser search; a web result is kept only when its URL appears in the search tool's evidence. Web searches receive only role, country and skill names, never CV text. Practice answers need separate consent, and feedback outlines use placeholders such as [your result] instead of inventing facts. Any provider failure falls back to clearly labelled curated or deterministic output. These checks do not prove semantic factuality: an earlier recorded rewrite added unsupported real-time sales monitoring, so human factual review remains necessary.")
    steps = [
        ("01", "CV + answers", "Opt-in + PII redaction\nCountry + target role\n3–5 follow-up questions\nAnswers reshape skills"),
        ("02", "Grounding", "CV + answer evidence\nFixed 2023 demand data\nAllowed skill vocabulary\nMatch % = skill overlap"),
        ("03", "AI + sources", f"{e['model']}\nGroq browser search\nHimalayas jobs API\nStructured JSON output"),
        ("04", "Checked output", "Quote + rewrite checks\nWeb results kept only\nif URL is in evidence\nNo-new-facts feedback")]
    for i, (n, heading, detail) in enumerate(steps):
        x = 76 + i * 370
        box(s, x, 379, 344, 270, "203A52", 12)
        text(s, n, x + 25, 396, 285, 56, 40, ORANGE, True, True)
        text(s, heading, x + 25, 458, 294, 42, 30, IVORY, True)
        text(s, detail, x + 25, 512, 305, 118, 22, PALE)
        if i < 3:
            text(s, "→", x + 344, 460, 28, 50, 27, ORANGE, True, align="center")
    box(s, 76, 681, 1452, 118, "29435B", 10)
    label(s, "PRIVACY BOUNDARY", 105, 701, 380)
    text(s, "Web search gets role, country + skill names — never CV text", 470, 698, 1030, 37, 26, IVORY, True)
    text(s, "CV to AI only with opt-in · practice answers need separate consent · provider failure → labelled fallback",
         105, 751, 1385, 30, 21, PALE)

    s = slide("05  /  Phone-first workflow", "A career next step, from a phone browser.",
              "Built for phone browsers — no app install. Editable panels below illustrate the new journey.",
              notes="These phone panels are code-native illustrations, not screenshots; the posting, match figure and feedback score are examples. Panel one: after the CV, role and consent, Njia asks follow-up questions with four options, Used it at work, Used it in a project or course, Still learning it, or Not yet, plus one free-text outcome question. Panel two: live jobs open to the learner's country or worldwide, each with a skill-match percentage, meaning the share of skills detected in the posting that the learner has, plus matched and missing skills and a link to apply on the source. It is not a hiring probability. Panel three: an interview question that candidates report on the web, with its source link. The learner types an answer under separate consent and gets a 1–5 score, a STAR checklist, strengths, improvements and a stronger outline with placeholders instead of invented facts; if the provider fails, a labelled deterministic checklist is shown. The seven-day checklist can also be sent to WhatsApp. The earlier release passed public desktop and 390 by 844 mobile-browser checks; its tested brief returned "
              f"{e['coverage']}% demand-weighted coverage over {e['sample_size']} Kenya Data Analyst postings.{jobs_note}")
    for i, (heading, subtitle) in enumerate([("01 / Give context", "CV + role + follow-up answers"),
                                             ("02 / Match live jobs", "Open to Kenya + skill match %"),
                                             ("03 / Practise", "Real questions + AI feedback")]):
        x = 78 + i * 495
        text(s, heading, x, 361, 455, 40, 30, INK, True)
        text(s, subtitle, x, 407, 455, 35, 23, MUTED)
        px = x + 46
        box(s, px, 462, 354, 343, NAVY, 24)
        box(s, px + 11, 474, 332, 318, WHITE, 16)
        box(s, px + 137, 480, 80, 6, NAVY, 3)
        text(s, "njia", px + 28, 496, 288, 37, 29, INK, True, True)
        if i == 0:
            label(s, "NJIA ASKS · 2 OF 4", px + 28, 545, 280, MUTED)
            text(s, "How have you used SQL?", px + 28, 577, 296, 30, 21, INK, True)
            for j, option in enumerate(["Used it at work", "Used it in a project or course",
                                        "Still learning it", "Not yet"]):
                oy = 617 + j * 40
                box(s, px + 26, oy, 302, 33, ORANGE if j == 1 else IVORY, 6)
                text(s, option, px + 38, oy + 5, 284, 24, 17, NAVY if j == 1 else INK, j == 1)
        elif i == 1:
            label(s, "LIVE · OPEN TO KENYA", px + 28, 545, 295, MUTED)
            text(s, "Data Analyst · Remote", px + 28, 577, 296, 30, 22, INK, True)
            text(s, "Example posting · worldwide", px + 28, 610, 296, 25, 17, MUTED)
            box(s, px + 26, 643, 302, 88, LIGHT, 6)
            text(s, "Skill match 67%", px + 39, 651, 280, 28, 22, INK, True)
            text(s, "Matched: SQL, Excel\nMissing: Python", px + 39, 683, 280, 44, 18, MUTED)
            box(s, px + 26, 742, 302, 32, ORANGE, 5)
            text(s, "Apply on source →", px + 30, 746, 294, 25, 18, NAVY, True, align="center")
        else:
            label(s, "FROM THE WEB · SOURCED", px + 28, 545, 295, MUTED)
            text(s, "“How do you handle\nmissing data?”", px + 28, 577, 296, 54, 21, INK, True)
            text(s, "Reported question · source link", px + 28, 633, 296, 24, 16, MUTED)
            box(s, px + 26, 666, 302, 108, LIGHT, 6)
            text(s, "Score 3 / 5 · STAR check", px + 38, 675, 283, 26, 19, INK, True)
            text(s, "Missing: the Result\nAdd: “…led to [your result]”", px + 38, 707, 283, 48, 18, INK)
    text(s, f"Panels illustrative. Skill match % = share of a posting's skills you have, not a hiring probability. Tested brief: {e['coverage']}% coverage, n = {e['sample_size']}.",
         76, 811, 1440, 24, 16, MUTED)

    s = slide("06  /  Evidence & honest limits", "A working product. Evidence you can inspect.",
              "Real public Groq calls · synthetic CVs · no mocked responses in the public checks.",
              notes=f"Main implementation owner confirms {bt} passing backend unit-test methods. Separately, {acc} public acceptance assertions and {rec} recording assertions passed, totaling {total} public assertions. No mocked responses. Desktop, mobile and recording used {e['provider']} {e['model']}. Observed elapsed times were {seconds('desktop')} desktop, {seconds('mobile')} mobile and {seconds('recording_with_pdf')} recording including PDF processing; these are individual observations, not a latency benchmark. The runs recorded {e['javascript_errors']} JavaScript errors. The sample used {e['sample_size']} Kenya Data Analyst postings and returned {e['coverage']}% demand-weighted coverage, not hiring odds. New today: " + "; ".join(new_feature_parts()) + ". Skill-match percentages measure skill overlap and are not a hiring probability. Semantic factuality is not guaranteed: the recorded rewrite added unsupported real-time sales monitoring, despite number/tool checks, and rewrites that use the learner's answers are labelled for verification. Human review is required. No real-user adoption or employment impact is claimed.")
    box(s, 76, 360, 666, 225, NAVY, 12)
    label(s, "PASSED / TWO DISTINCT TEST LAYERS", 106, 381, 600, PALE)
    text(s, str(bt), 106, 418, 300, 98, 82, IVORY, True, True)
    text(s, str(total), 435, 418, 280, 98, 82, ORANGE, True, True)
    text(s, "backend tests", 110, 519, 290, 36, 25, PALE)
    text(s, "public assertions", 439, 518, 290, 33, 24, PALE)
    text(s, f"{acc} acceptance + {rec} recording", 439, 553, 290, 25, 18, PALE)
    box(s, 780, 360, 748, 225, LIGHT, 12)
    label(s, "LIVE GROQ BRIEF / OBSERVED RESPONSE TIMES", 812, 381, 690, INK)
    text(s, f"{seconds('desktop')} desktop  /  {seconds('mobile')} mobile\n{seconds('recording_with_pdf')} recording, including PDF processing\n"
            f"{e['javascript_errors']} JavaScript errors across these runs", 812, 424, 690, 100, 27, INK)
    text(s, "Individual observations, not a performance benchmark.", 812, 540, 690, 27, 20, MUTED)
    box(s, 76, 603, 1452, 80, WHITE, 10)
    box(s, 76, 603, 6, 80, ORANGE)
    label(s, "NEW TODAY / FOLLOW-UPS · LIVE JOBS · WEB INTERVIEW QUESTIONS · ANSWER FEEDBACK", 106, 613, 1400, INK)
    text(s, "  ·  ".join(new_feature_parts()), 106, 645, 1400, 28, 21, MUTED)
    for x, head, body in [
        (76, "Data + match %", "2023 sample sets the priorities.\nMatch % is not a hiring probability."),
        (573, "Human factual review", "Numbers/tools checked; unsupported\nclaims can still slip through."),
        (1070, "Impact", "No measured employment outcomes\nor real-user adoption claimed.")]:
        text(s, head, x, 702, 470, 40, 28, INK, True, True)
        text(s, body, x, 746, 470, 52, 21, MUTED)

    s = slide("07  /  Adoption plan & award fit", "Start small. Measure a useful next step.",
              "Proposed pilot with Kenyan learners and a training partner — recruitment has not started.",
              notes="This is a proposed pilot, not adoption or a signed partnership. We propose recruiting ten Kenyan learners with one training partner, reviewing outputs with a facilitator and following up after seven days. Measure completion, factual rewrite acceptance, whether follow-up answers were reflected, relevance of the live jobs shown, usefulness of answer feedback, saved practice artifacts and mobile friction; do not present target outcomes as achieved. Click Mobile fit is the practical mobile-first Kenyan learner journey, with Dhruzzz participating as Kenya ONLINE. Published Kenya-only scope aligns with the team context; final award eligibility is determined by organizers. Brightest fits skills and employability, and Artefact fits visible use of local data and AI. No award win or organizer endorsement is claimed.")
    for i, (tag, title, desc) in enumerate([
        ("RECRUIT", "10 learners + 1 partner", "Invite Kenyan career starters.\nObserve the mobile CV journey."),
        ("REVIEW", "Human-check the advice", "Check that rewrites keep the facts\nand job matches are relevant."),
        ("FOLLOW UP", "Return after 7 days", "Count practice answers and\napplications started. Ask what\nhelped and what blocked progress.")]):
        x = 76 + i * 497
        label(s, tag, x, 387, 455, MUTED)
        box(s, x, 433, 452, 4, ORANGE)
        text(s, title, x, 465, 454, 78, 33, INK, True, True)
        text(s, desc, x, 560, 465, 97, 24, MUTED)
    box(s, 76, 687, 1452, 117, NAVY, 12)
    label(s, "CLICK MOBILE / KENYA ONLINE", 108, 708, 570)
    text(s, "Mobile-first coaching for Kenyan learners.", 660, 706, 840, 38, 28, IVORY, True)
    text(s, "Award fit: practical phone-browser access · also aligned with Brightest (employability) and Artefact (Data & AI).",
         108, 757, 1400, 28, 20, PALE)

    demo_note = ("It shows the earlier brief flow and predates today's follow-up questions, live jobs, web interview "
                 "questions and answer feedback. " if e["demo_predates_new_features"] else "")
    s = slide("08  /  Team Dhruzzz", "Turn experience into a clearer path.", dark=True,
              notes=f"Close with the live product and source repository. Team: Eeshan Vaghjiani, Bhavin Mepani and Dhruvin Bhudia; Kenya ONLINE. Our line: {TAGLINE} The published demo-v1 recording is exactly {e['video_seconds']} seconds, normal speed, with captions. {demo_note}Slide 8 links to it: {DEMO_URL}. Ask for a Kenyan learner or training partner to run the proposed pilot.")
    text(s, "Njia", 76, 281, 680, 133, 109, IVORY, True, True)
    text(s, TAGLINE, 80, 425, 1415, 69, 39, ORANGE, serif=True)
    label(s, "KENYA · ONLINE", 80, 532, 600, PALE)
    text(s, "Eeshan Vaghjiani\nBhavin Mepani\nDhruvin Bhudia", 80, 582, 620, 155, 34, IVORY)
    box(s, 794, 529, 734, 269, "203A52", 12)
    for y, tag, display, url in [
        (552, "LIVE", "gomycode-2026.vercel.app", "https://gomycode-2026.vercel.app"),
        (626, "SOURCE", "github.com/Eeshan-Vaghjiani/njia", "https://github.com/Eeshan-Vaghjiani/njia"),
        (700, f"{e['video_seconds']}-SECOND DEMO", "njia-demo-90s.webm · published", DEMO_URL)]:
        label(s, tag, 827, y, 630)
        text(s, display, 827, y + 30, 669, 38, 22 if url == DEMO_URL else 24, IVORY, link=url)
    text(s, "Pilot invitation: Kenyan learners + training partners", 80, 768, 700, 37, 24, PALE)
    return dict(total_postings=len(rows), countries=countries, kenya_postings=len(ke),
                kenya_data_analyst_postings=len(da), chart=dict(demand))


def make_pptx():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    prs.core_properties.title = "Njia — Team Dhruzzz"
    prs.core_properties.subject = "Ask. Match. Practise. Evidence-grounded career coaching for Kenyan learners"
    prs.core_properties.author = "Team Dhruzzz"
    prs.core_properties.keywords = "Njia, Kenya, Dhruzzz, career guidance, live jobs, interview practice, historical demand"
    for s in SLIDES:
        page = prs.slides.add_slide(prs.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(s["bg"])
        for e in s["items"]:
            geometry = [Inches(e[k] / 100) for k in ("x", "y", "w", "h")]
            if e["kind"] == "box":
                sh = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE, *geometry)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
                sh.line.fill.background()
                if e["radius"]:
                    sh.adjustments[0] = min(.3, e["radius"] / min(e["w"], e["h"]))
            else:
                sh = page.shapes.add_textbox(*geometry)
                tf = sh.text_frame
                tf.clear()
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                for index, line in enumerate(e["text"].split("\n")):
                    p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
                    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[e["align"]]
                    p.line_spacing = 1.12
                    p.space_before = p.space_after = Pt(0)
                    r = p.add_run()
                    r.text = line
                    r.font.name = "Georgia" if e["serif"] else "Segoe UI"
                    r.font.size = Pt(e["size"] * .72)
                    r.font.bold = e["bold"]
                    r.font.color.rgb = RGBColor.from_string(e["color"])
                    if e["link"]:
                        r.hyperlink.address = e["link"]
        page.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(OUT / f"{STEM}.pptx")


def make_html():
    fragments = []
    for index, s in enumerate(SLIDES):
        items = []
        for e in s["items"]:
            style = f'left:{e["x"]}px;top:{e["y"]}px;width:{e["w"]}px;height:{e["h"]}px;'
            if e["kind"] == "box":
                style += f'background:#{e["fill"]};border-radius:{e["radius"]}px;'
                items.append(f'<div class="shape" style="{style}"></div>')
            else:
                style += (f'font-size:{e["size"]}px;color:#{e["color"]};'
                          f'font-weight:{700 if e["bold"] else 400};text-align:{e["align"]};'
                          f'font-family:{SERIF if e["serif"] else SANS};')
                body = html.escape(e["text"]).replace("\n", "<br>")
                if e["link"]:
                    body = f'<a href="{html.escape(e["link"])}">{body}</a>'
                items.append(f'<div class="text" style="{style}">{body}</div>')
        fragments.append(f'<section class="slide" aria-label="Slide {index+1}" style="background:#{s["bg"]}">{"".join(items)}</section>')
    document = '''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=1600"><title>Njia — Team Dhruzzz</title>
<style>@page{size:1600px 900px;margin:0}*{box-sizing:border-box}html,body{margin:0;padding:0}
body{background:#c9cccf}.slide{position:relative;width:1600px;height:900px;overflow:hidden;
break-after:page;page-break-after:always}.slide:last-child{break-after:auto;page-break-after:auto}
.shape,.text{position:absolute}.text{line-height:1.12;white-space:normal;overflow:visible}
a{color:inherit;text-decoration:none} @media print{body{background:white}
*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}@media screen{.slide{margin:0 auto 24px}}</style>
</head><body>''' + "".join(fragments) + "</body></html>"
    (OUT / f"{STEM}.html").write_text(document, encoding="utf-8")


def render():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1600, "height": 900}, device_scale_factor=1)
        page.goto((OUT / f"{STEM}.html").as_uri(), wait_until="load")
        page.evaluate("document.fonts.ready")
        overflow = page.locator(".text").evaluate_all("""els => els.filter(e =>
            e.scrollHeight > e.clientHeight + 2 || e.scrollWidth > e.clientWidth + 2)
            .map(e => ({text:e.innerText, actual:[e.scrollWidth,e.scrollHeight],
            box:[e.clientWidth,e.clientHeight]}))""")
        if overflow:
            raise RuntimeError("Text overflow: " + json.dumps(overflow, ensure_ascii=False))
        for index, section in enumerate(page.locator(".slide").all()):
            section.screenshot(path=str(OUT / f"slide-{index+1:02d}.png"))
        page.emulate_media(media="print")
        page.pdf(path=str(OUT / f"{STEM}.pdf"), prefer_css_page_size=True,
                 print_background=True, display_header_footer=False)
        browser.close()
    print(f"Rendered 8 slide previews and PDF in {OUT}; HTML text-overflow check passed.")


def required_values():
    e = EVIDENCE
    values = ["Dhruzzz", "Eeshan Vaghjiani", "Bhavin Mepani", "Dhruvin Bhudia", "18,371", "1,326",
              str(e["sample_size"]), str(e["backend_tests"]), str(public_total()),
              f'{e["acceptance_assertions"]} acceptance + {e["recording_assertions"]} recording',
              f'{e["coverage"]}%', seconds("desktop"), seconds("mobile"), seconds("recording_with_pdf"),
              TAGLINE, "not a hiring probability", "never CV text", "Uses your answer — verify",
              "Which jobs can I apply for now?", "Am I ready for the interview?"]
    if e["live_jobs_example"]:
        values.append(f'{e["live_jobs_example"]["count"]:,} remote')
    if e["web_questions_example"]:
        values.append(f'{e["web_questions_example"]["count"]} source-linked')
    if e["new_feature_tests"]:
        values.append(f'{e["new_feature_tests"]} new-feature tests')
    return values


def validate():
    import fitz
    from PIL import Image, ImageDraw
    from pptx import Presentation
    prs = Presentation(OUT / f"{STEM}.pptx")
    pdf = fitz.open(OUT / f"{STEM}.pdf")
    assert len(prs.slides) == len(pdf) == 8
    slide_text = "\n".join(shape.text for s in prs.slides for shape in s.shapes if shape.has_text_frame)
    pdf_text = "\n".join(p.get_text() for p in pdf)
    flat_pdf_text = re.sub(r"\s+", " ", pdf_text)
    for value in required_values():
        assert value.casefold() in slide_text.casefold(), value
        assert value.casefold() in flat_pdf_text.casefold(), value
    notes_text = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
    supporting_text = "\n".join((OUT / name).read_text(encoding="utf-8") for name in
                                ["speaker-notes.md", "README.md", "production-evidence.json"])
    filled = json.dumps(EVIDENCE)
    for content in [slide_text, pdf_text, notes_text, supporting_text]:
        if not re.search(r"\b72\b", filled):
            assert not re.search(r"\b72\b", content), "Stale backend test count"
        assert "pending" not in content.casefold(), "Stale release status"
    for content in [slide_text, pdf_text, notes_text]:
        assert not re.search(r"\bNone\b", content), "Unfilled EVIDENCE value rendered as None"
    evidence = json.loads((OUT / "production-evidence.json").read_text(encoding="utf-8"))
    assert evidence["backend_tests_passed"] == EVIDENCE["backend_tests"]
    assert (evidence["public_acceptance_assertions"], evidence["recording_assertions"], evidence["public_total"]) == (
        EVIDENCE["acceptance_assertions"], EVIDENCE["recording_assertions"], public_total())
    assert evidence["public_acceptance_assertions"] + evidence["recording_assertions"] == evidence["public_total"]
    assert "human factual review required" in evidence["factual_limit"].casefold()
    assert evidence["unverified_new_feature_fields"] == [k for k in NEW_FEATURE_KEYS if EVIDENCE[k] is None]
    assert "Semantic factuality is not guaranteed" in notes_text
    assert "not a hiring probability" in notes_text
    assert "never CV text" in notes_text
    pptx_links = [r.hyperlink.address for shape in prs.slides[7].shapes if shape.has_text_frame
                  for p in shape.text_frame.paragraphs for r in p.runs]
    assert DEMO_URL in pptx_links
    assert DEMO_URL in [link.get("uri") for link in pdf[7].get_links()]
    html_text = (OUT / f"{STEM}.html").read_text(encoding="utf-8")
    assert f'href="{DEMO_URL}"' in html_text
    assert "pending" not in html_text.casefold()
    for slide_index, s in enumerate(prs.slides):
        assert s.notes_slide.notes_text_frame.text.strip()
        for shape in s.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
    for page in pdf:
        assert len(page.get_text()) > 200
        assert math.isclose(page.rect.width / page.rect.height, 16/9, abs_tol=.001)
    sheet = Image.new("RGB", (1280, 4 * 390), "#D4DEE6")
    draw = ImageDraw.Draw(sheet)
    for i, page in enumerate(pdf):
        pix = page.get_pixmap(matrix=fitz.Matrix(.5, .5), alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img.thumbnail((624, 351))
        x, y = (i % 2) * 640 + 8, (i // 2) * 390 + 27
        sheet.paste(img, (x, y))
        draw.text((x, y - 20), f"SLIDE {i+1:02d}", fill="#14283F")
    sheet.save(OUT / "contact-sheet.png")
    bt, acc, rec, total = (EVIDENCE["backend_tests"], EVIDENCE["acceptance_assertions"],
                           EVIDENCE["recording_assertions"], public_total())
    result = dict(pptx_slides=len(prs.slides), pdf_pages=len(pdf), aspect_ratio="16:9",
                  pptx_bytes=(OUT / f"{STEM}.pptx").stat().st_size,
                  pdf_bytes=(OUT / f"{STEM}.pdf").stat().st_size,
                  checks=["Slide/page count", "Required content in PPTX and PDF", "PPTX shape bounds",
                          "Speaker notes on all slides", "PDF searchable text", "PDF aspect ratio",
                          f"{bt} backend tests; {total} public assertions = {acc} + {rec}",
                          "No stale count, unfilled value or stale release text",
                          "Published demo URL in PPTX, PDF and HTML", "Human factual-review caveat retained",
                          "Not-a-hiring-probability and web-search privacy boundary retained"],
                  unverified_new_feature_fields=evidence["unverified_new_feature_fields"],
                  advisor_validation=f"{bt} backend tests confirmed by main owner; {acc} public acceptance + {rec} recording assertions; human factual review required")
    (OUT / "validation.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


def write_supporting(evidence):
    e = EVIDENCE
    (OUT / "data-evidence.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    (OUT / "speaker-notes.md").write_text("# Njia — speaker notes\n\n" + "\n\n".join(
        f'## {i+1}. {s["title"].replace(chr(10), " ")}\n\n{s["notes"]}' for i, s in enumerate(SLIDES)), encoding="utf-8")
    new_features = "; ".join(new_feature_parts())
    (OUT / "README.md").write_text(f'''# Njia — Team Dhruzzz

Local eight-slide final deck in navy, ivory and orange: **{TAGLINE}** The PPTX contains
editable text, shapes, chart bars and phone-flow illustrations; all slides include speaker
notes. The PDF is printed from the matching local HTML with Chromium/Playwright. No network
assets or model calls are required to build these files.

## Deliverables
- `Njia-Dhruzzz.pptx` — editable presentation, 8 slides.
- `Njia-Dhruzzz.pdf` — searchable PDF, 8 pages.
- `Njia-Dhruzzz.html` — matching printable HTML.
- `speaker-notes.md` — presentation narration and supporting qualifications.
- `contact-sheet.png` and `slide-01.png` … `slide-08.png` — visual previews.
- `data-evidence.json` — locally recalculated denominators and chart values.
- `production-evidence.json` — every tested or observed figure used in the deck.
- `validation.json` — file sizes and structural validation results.

## Product story
Njia asks before it advises: after the CV it asks a few follow-up questions whose answers
reshape suggested skills and the brief. It matches the learner to live jobs they can apply
for, with a skill-match % (skill overlap, not a hiring probability), and supports interview
practice with source-linked questions from the web and AI feedback on typed answers. Web
searches receive only role, country and skill names — never CV text. Provider failures fall
back to labelled curated or deterministic output.

## Current production evidence
{e["backend_tests"]} backend test methods passed ({e["backend_count_source"].lower()}), distinct from
{public_total()} public assertions: {e["acceptance_assertions"]} acceptance + {e["recording_assertions"]} recording.
Real {e["provider"]} / {e["model"]}: {seconds("desktop")} desktop, {seconds("mobile")} mobile, {seconds("recording_with_pdf")} recording including PDF.
Individual observations, not a performance benchmark. {e["javascript_errors"]} JavaScript errors in these runs.
Sample: {e["sample_size"]} historical Kenya Data Analyst postings, {e["coverage"]}% demand-weighted coverage, not hiring odds.
New features: {new_features}.
See `production-evidence.json` for the precise count breakdown and unverified fields.

CV rewrites checked for numbers/tools, still require human factual review. The recording
contains unsupported “real-time sales monitoring” in a rewrite. The checks do not guarantee
semantic factuality. Rewrites that use the learner's own answers are labelled for verification.
The deck's before/after example is illustrative and hand-authored.

## Published demo
The published `demo-v1` recording is exactly {e["video_seconds"]} seconds. Slide 8 links to it:
{DEMO_URL}

Illustrative persona, CV output and phone panels are explicitly labeled. The pilot is a plan,
not adoption. The priority dataset is a 2023 historical sample, not current vacancies. CV
transfer to AI requires opt-in and basic, best-effort PII redaction; there is no application
database persistence claim beyond the application itself.

## Rebuild locally (from project root)
```bash
.venv/bin/python scripts/build_final_deck.py --all
.venv/bin/python scripts/build_final_deck.py --all --out /tmp/opencode/deck
```
On Windows use `.venv\\Scripts\\python.exe`. `--all` builds the PPTX/HTML, renders the PDF and
PNG previews with Playwright, then validates. Update `EVIDENCE` in
`scripts/build_final_deck.py` before rebuilding. The builder writes only inside `--out`
(default `presentation/`).
''', encoding="utf-8")


def main():
    global OUT
    parser = argparse.ArgumentParser()
    parser.add_argument("--render-only", action="store_true")
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("--all", action="store_true", help="build, render PDF/PNGs, then validate")
    parser.add_argument("--out", default=str(ROOT / "presentation"), help="output directory (default: presentation/)")
    args = parser.parse_args()
    OUT = Path(args.out).expanduser().resolve()
    OUT.mkdir(parents=True, exist_ok=True)
    if args.render_only:
        render()
        return
    if args.validate:
        validate()
        return
    evidence = build_content()
    assert len(SLIDES) == 8
    make_pptx()
    make_html()
    write_supporting(evidence)
    print(f"Built editable 8-slide PPTX + matching HTML in {OUT}.")
    print(json.dumps(evidence, indent=2))
    if args.all:
        render()
        validate()


if __name__ == "__main__":
    main()
